"""SSB form contracts from the public API through original Word cells."""

from io import BytesIO

from django.contrib.auth import get_user_model
from docx import Document
from rest_framework.test import APITestCase

from .form_processes.catalog import FORM_TEMPLATES
from .form_processes.models import FormProcessRecord


class SsbFormTests(APITestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(username="ssb-test", is_active=True)
        self.client.force_authenticate(self.user)

    def payload(self):
        return {
            "template_code": "pr_qua_20_104E",
            "record_number": "SSB-001",
            "title": "SSB test başvurusu",
            "status": "approved",
            "data": {
                "contract_number": "SOZ-001",
                "applicant": "Test & Geliştirme <Birimi>",
                "application_number": "UI-001",
                "aircraft_nationality": "Türkiye",
                "aircraft_id_mark": "TEST-01",
                "aircraft_owner": "Test kuruluşu",
                "aircraft_manufacturer": "Test üreticisi",
                "aircraft_model": "Model A",
                "serial_number": "SN-001",
                "intended_flight_date": "2026-09-25",
                "flight_duration": "2",
                "purpose_of_flight": ["option_1"],
                "purpose_scope": "Uçuş testi\nİkinci aşama",
                "aircraft_configuration": "Konfigürasyon A",
                "conditions_restrictions": "Gündüz uçuşu",
                "substantiations": "Rapor R-001",
                "issue_date": "2026-09-22",
                "approver_name": "Yüklenici Yetkilisi",
                "board_chair_name": "Kurul Başkanı",
                "board_psk_name": "PSK Yetkilisi",
                "board_members": [{"name": f"Üye {i}"} for i in range(1, 5)],
                "valid_from": "2026-09-23",
                "valid_until": "2026-09-30",
                "permit_issue_date": "2026-09-23",
                "is_recommendation": "yes",
                "flight_test_plan_number": "FTP-01",
                "permit_lifecycle_status": "draft",
            },
        }

    def test_catalog_metadata_matches_named_fields_for_every_form(self):
        response = self.client.get("/api/form-processes/templates/")
        self.assertEqual(response.status_code, 200)
        catalog = {t["code"]: t for p in response.data for t in p["templates"]}
        for definition in FORM_TEMPLATES:
            with self.subTest(template=definition.code):
                entry = catalog[definition.code]
                self.assertEqual(
                    entry["form_number"].replace(".", "_").lower(), definition.code.lower()
                )
                self.assertNotEqual(entry["title"], entry["form_number"])
        self.assertEqual(catalog["pr_qua_20_104E"]["title"], "SSB Özel Uçuş İzni Tavsiyesi Formu")

    def test_create_edit_approve_download_fills_original_cells(self):
        payload = self.payload()
        created = self.client.post(
            "/api/form-processes/", {**payload, "status": "draft", "data": {}}, format="json"
        )
        self.assertEqual(created.status_code, 201, created.data)
        url = f"/api/form-processes/{created.data['id']}/"
        saved = self.client.patch(
            url, {"data": payload["data"], "status": "approved"}, format="json"
        )
        self.assertEqual(saved.status_code, 200, saved.data)
        self.assertEqual(saved.data["form_number"], "PR.QUA.20.104E")
        response = self.client.get(saved.data["generated_document_url"])
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["X-Content-Type-Options"], "nosniff")
        document = Document(BytesIO(b"".join(response.streaming_content)))
        table = document.tables[0]
        expected = {
            (0, 1): "SOZ-001",
            (1, 1): payload["data"]["applicant"],
            (2, 1): "UI-001",
            (3, 1): "Türkiye / TEST-01",
            (4, 1): "Test kuruluşu",
            (5, 1): "Test üreticisi / Model A",
            (6, 1): "SN-001",
            (7, 1): "25.09.2026 / 2 saat",
            (8, 0): "Uçuş testi\nİkinci aşama",
            (9, 0): "Konfigürasyon A",
            (10, 0): "Gündüz uçuşu",
            (11, 0): "Rapor R-001",
            (13, 1): "22.09.2026",
            (14, 1): "Yüklenici Yetkilisi",
            (15, 2): "Kurul Başkanı",
            (16, 2): "PSK Yetkilisi",
            (21, 1): "23.09.2026 – 30.09.2026",
            (22, 1): "23.09.2026",
        }
        for (row, column), value in expected.items():
            with self.subTest(row=row, column=column):
                self.assertIn(value, table.cell(row, column).text)
        for index in range(4):
            self.assertEqual(table.cell(17 + index, 2).text, f"Üye {index + 1}")
        self.assertIn("Geliştirme", table.cell(8, 0).text)
        self.assertIn("Hava aracı, tanımlanmış", table.cell(12, 0).text)
        for row in range(15, 21):
            self.assertEqual(table.cell(row, 3).text, "İmza")
            self.assertEqual(table.cell(row, 4).text, "Paraf")
        self.assertNotIn(
            "{{", "\n".join(c.text for t in document.tables for r in t.rows for c in r.cells)
        )

    def test_invalid_values_are_rejected_without_writes(self):
        cases = [
            ("valid_until", "2026-09-01"),
            ("flight_duration", "0"),
            ("flight_duration", "-2"),
            ("flight_duration", "1.5"),
            ("flight_duration", "²"),
            ("flight_duration", "abc"),
            ("flight_duration", "9" * 5000),
            ("intended_flight_date", "2026-09-22"),
            ("intended_flight_date", "2026-10-01"),
            ("permit_issue_date", "2026-02-30"),
            ("board_members", [{"name": "Üye"}] * 5),
            ("board_members", [{"name": 42}]),
            ("purpose_of_flight", ["unknown"]),
            ("purpose_of_flight", [{}]),
            ("purpose_of_flight", [["option_1"]]),
            ("permit_issue_date", "20260923"),
            ("applicant", ""),
        ]
        for field, value in cases:
            with self.subTest(field=field, value=value):
                payload = self.payload()
                payload["data"][field] = value
                response = self.client.post("/api/form-processes/", payload, format="json")
                self.assertEqual(response.status_code, 400, response.data)
                self.assertIn(field, response.data)
                self.assertFalse(FormProcessRecord.objects.exists())

    def test_invalid_update_preserves_saved_record(self):
        payload = self.payload()
        created = self.client.post("/api/form-processes/", payload, format="json")
        self.assertEqual(created.status_code, 201, created.data)
        previous = FormProcessRecord.objects.get().data
        payload["data"]["valid_until"] = "2026-09-01"
        response = self.client.patch(
            f"/api/form-processes/{created.data['id']}/", {"data": payload["data"]}, format="json"
        )
        self.assertEqual(response.status_code, 400)
        self.assertEqual(FormProcessRecord.objects.get().data, previous)

    def test_flight_on_either_validity_boundary_and_legacy_payload_are_accepted(self):
        for index, flight_date in enumerate(["2026-09-23", "2026-09-30"]):
            payload = self.payload()
            payload["record_number"] = f"SSB-BOUNDARY-{index}"
            payload["data"]["intended_flight_date"] = flight_date
            for key in ("board_chair_name", "board_psk_name", "board_members", "permit_issue_date"):
                del payload["data"][key]
            response = self.client.post("/api/form-processes/", payload, format="json")
            self.assertEqual(response.status_code, 201, response.data)
            download = self.client.get(response.data["generated_document_url"])
            self.assertEqual(download.status_code, 200)
            download.close()

    def test_optional_cells_and_draft_remain_blank(self):
        response = self.client.post(
            "/api/form-processes/", {**self.payload(), "status": "draft", "data": {}}, format="json"
        )
        self.assertEqual(response.status_code, 201, response.data)
        download = self.client.get(response.data["generated_document_url"])
        document = Document(BytesIO(b"".join(download.streaming_content)))
        for row, column in [(0, 1), (7, 1), (13, 1), (15, 2), (20, 2), (21, 1), (22, 1)]:
            self.assertEqual(document.tables[0].cell(row, column).text, "")
        rejected = self.client.patch(
            f"/api/form-processes/{response.data['id']}/", {"status": "approved"}, format="json"
        )
        self.assertEqual(rejected.status_code, 400)
        self.assertEqual(FormProcessRecord.objects.get().status, "draft")

    def test_shared_access_and_inactive_or_anonymous_denial(self):
        response = self.client.post("/api/form-processes/", self.payload(), format="json")
        self.assertEqual(response.status_code, 201, response.data)
        url = response.data["generated_document_url"]
        for active, staff, expected in [(True, False, 200), (True, True, 200), (False, False, 403)]:
            user = get_user_model().objects.create_user(
                username=f"reader-{active}-{staff}", is_active=active, is_staff=staff
            )
            self.client.force_authenticate(user)
            download = self.client.get(url)
            self.assertEqual(download.status_code, expected)
            download.close()
        self.client.force_authenticate(None)
        self.assertEqual(self.client.get(url).status_code, 403)
